"""Fill the empty right-hand (transliteration) column of a Word file of duas:
one table per dua, Arabic in the left cell. Used for Dad's dua file.

    python3 dua_docx.py in.docx out.docx [first last]   # tables first..last-1

dua_fixes.json lists corrections to the source Arabic ("fixes": table, old,
new, note) and to the transliteration where the program gets a correct
Arabic word wrong ("translit": table, got, want). Corrected Arabic is
written into the Arabic column, and the changed words are highlighted yellow
in both columns so they can be checked against the book.
"""
import copy, difflib, json, os, re, sys
import docx
from docx.enum.text import WD_COLOR_INDEX
import dua

HERE = os.path.dirname(os.path.abspath(__file__))
FIXES = json.load(open(os.path.join(HERE, 'dua_fixes.json'), encoding='utf-8'))


def canon(t):
    # one order for stacked marks (shadda + vowel) so the fixes match
    return re.sub('[ً-ٰٗ]+', lambda m: ''.join(sorted(m.group(0))), t)


def fixed(i, ar):
    ar = canon(ar)
    for f in FIXES['fixes']:
        if f['table'] == i and f.get('sure', True):
            old, new = canon(f['old']), canon(f['new'])
            assert ar.count(old) == 1, (i, f['old'], ar.count(old))
            ar = ar.replace(old, new)
    return ar


def translit(i, ar):
    text = dua.transliterate(ar)
    for f in FIXES['translit']:
        if f['table'] == i:
            assert text.count(f['got']) == 1, (i, f['got'], text.count(f['got']))
            text = text.replace(f['got'], f['want'])
    return text


def changed(before, after):
    """-> [(chunk, is_changed)] covering `after`, by words."""
    a, b = before.split(), after.split()
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes():
        if j2 > j1:
            out.append((' '.join(b[j1:j2]), op != 'equal'))
    return out


def fill(p, chunks, rpr=None):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for k, (txt, hl) in enumerate(chunks):
        r = p.add_run(txt + (' ' if k + 1 < len(chunks) else ''))
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        if hl:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def arabic_chunks(i, orig, new):
    """Words of the corrected Arabic as [(text, colour)]: changed words yellow,
    words of an unsure suggestion (not applied) green, the rest as typed."""
    o, n = orig.split(), new.split()
    words = []
    sm = difflib.SequenceMatcher(a=[canon(w) for w in o], b=n, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == 'equal':
            words += [[w, None] for w in o[i1:i2]]
        else:
            words += [[w, 'Y'] for w in n[j1:j2]]
    # character offsets of each word in `new`
    pos, starts = 0, []
    for w in n:
        pos = new.index(w, pos)
        starts.append(pos)
        pos += len(w)
    for f in FIXES['fixes']:
        if f['table'] == i and not f.get('sure', True):
            at = new.find(canon(f['old']))
            if at < 0:
                continue
            for k, st in enumerate(starts):
                if st < at + len(canon(f['old'])) and st + len(n[k]) > at and words[k][1] is None:
                    words[k][1] = 'G'
    chunks = []
    for w, c in words:
        if chunks and chunks[-1][1] == c:
            chunks[-1][0] += ' ' + w
        else:
            chunks.append([w, c])
    return chunks


def fill(p, chunks, rpr=None):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for k, (txt, c) in enumerate(chunks):
        r = p.add_run(txt + (' ' if k + 1 < len(chunks) else ''))
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        if c:
            r.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN if c == 'G' else WD_COLOR_INDEX.YELLOW


def corrections_list(d):
    """Append every correction at the end of the document."""
    d.add_page_break()
    d.add_heading('Corrections to the Arabic', level=1)
    d.add_paragraph('Yellow in the duas: corrected (listed below). Green: a possible '
                    'correction that was not made; please check it against the book.')
    for sure, title in ((True, 'Corrected'), (False, 'Please check (not changed)')):
        d.add_heading(title, level=2)
        rows = [f for f in FIXES['fixes'] if f.get('sure', True) == sure]
        t = d.add_table(rows=1, cols=4)
        if 'Table Grid' in [st.name for st in d.styles]:
            t.style = 'Table Grid'
        for c, h in zip(t.rows[0].cells, ('Dua', 'Was', 'Now' if sure else 'Suggested', 'Note')):
            c.text = h
        for f in sorted(rows, key=lambda f: f['table']):
            r = t.add_row().cells
            r[0].text = str(f['table'] + 1)
            r[1].text, r[2].text, r[3].text = f['old'], f['new'], f['note']
    if FIXES.get('questions'):
        d.add_heading('Questions', level=2)
        for q in FIXES['questions']:
            d.add_paragraph('• ' + q)


def main(src, out, lo=0, hi=None):
    d = docx.Document(src)
    hi = len(d.tables) if hi is None else hi
    for i in range(lo, hi):
        ar_cell, tr_cell = d.tables[i].rows[0].cells
        orig = ar_cell.text
        new = fixed(i, orig)
        chunks = arabic_chunks(i, orig, new)
        if any(c for _, c in chunks):
            p = ar_cell.paragraphs[0]
            rpr = p.runs[0]._r.rPr if p.runs and p.runs[0]._r.rPr is not None else None
            fill(p, chunks, rpr)
        text = translit(i, new)
        fill(tr_cell.paragraphs[0], [(t, 'Y' if c else None) for t, c in changed(translit(i, canon(orig)), text)])
    if lo == 0 and hi == len(d.tables):
        corrections_list(d)
    d.save(out)


if __name__ == '__main__':
    a = sys.argv[1:]
    main(a[0], a[1], *map(int, a[2:4]))
